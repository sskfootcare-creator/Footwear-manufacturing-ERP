"""Enterprise CSV & Excel (XLSX) Export Service.

Provides CA-ready, audit-grade tabular exports for:
- Overdue & General Invoices
- Karigar Payroll Statements
- Customer Purchase Orders
- Client Balances & AR Ledgers
- Vendor Procurement & AP Ledgers
- Finished Goods & Materials Inventory
- QC Defects & Rework Log
"""

import io
import csv
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any
from fastapi.responses import StreamingResponse

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def generate_csv_bytes(headers: List[str], rows: List[List[Any]]) -> bytes:
    """Generate RFC 4180 compliant CSV bytes with UTF-8 BOM for seamless Excel compatibility."""
    output = io.StringIO()
    # Write UTF-8 BOM so Excel opens Hindi / special characters & numbers cleanly
    output.write("\uFEFF")
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for r in rows:
        writer.writerow([str(c) if c is not None else "" for c in r])
    return output.getvalue().encode("utf-8")


def generate_excel_workbook(
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    summary_kpis: Optional[Dict[str, Any]] = None,
    sheet_name: str = "ExportData",
    currency_col_indices: Optional[List[int]] = None,
    total_col_indices: Optional[List[int]] = None,
) -> bytes:
    """Generates an audit-grade Excel Workbook formatted with typography, styling,
    zebra striping, currency number formats, KPI cards, and total summary row.
    """
    if not HAS_OPENPYXL:
        # Fallback to CSV if openpyxl is somehow absent
        return generate_csv_bytes(headers, rows)

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Sheet1")[:31]

    # Typography & Colors
    font_family = "Segoe UI"
    title_font = Font(name=font_family, size=14, bold=True, color="0F172A")
    meta_font = Font(name=font_family, size=9, italic=True, color="64748B")
    tbl_hdr_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    tbl_hdr_fill = PatternFill("solid", fgColor="0F172A")  # Slate 900
    subhdr_fill = PatternFill("solid", fgColor="C27842")   # SSK Gold/Amber accent
    zebra_fill = PatternFill("solid", fgColor="F8FAFC")
    total_fill = PatternFill("solid", fgColor="E2E8F0")

    regular_font = Font(name=font_family, size=10)
    bold_font = Font(name=font_family, size=10, bold=True)
    
    red_status_fill = PatternFill("solid", fgColor="FEE2E2")
    red_status_font = Font(name=font_family, size=9, bold=True, color="991B1B")
    green_status_fill = PatternFill("solid", fgColor="DCFCE7")
    green_status_font = Font(name=font_family, size=9, bold=True, color="166534")
    amber_status_fill = PatternFill("solid", fgColor="FEF3C7")
    amber_status_font = Font(name=font_family, size=9, bold=True, color="92400E")

    thin_border_side = Side(style="thin", color="CBD5E1")
    double_bottom_side = Side(style="double", color="0F172A")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    total_border = Border(top=thin_border_side, bottom=double_bottom_side, left=thin_border_side, right=thin_border_side)

    curr_indices = set(currency_col_indices or [])
    tot_indices = set(total_col_indices or [])

    curr_row = 1

    # 1. Title Banner
    ws.cell(row=curr_row, column=1, value=title).font = title_font
    curr_row += 1

    now_str = datetime.now(timezone.utc).strftime("%d-%b-%Y %H:%M UTC")
    meta_text = f"SSK Footcare Manufacturing ERP · Generated: {now_str} · Total Records: {len(rows)}"
    ws.cell(row=curr_row, column=1, value=meta_text).font = meta_font
    curr_row += 2

    # 2. KPI Cards (if provided)
    if summary_kpis:
        kpi_col = 1
        for kpi_label, kpi_val in summary_kpis.items():
            lbl_cell = ws.cell(row=curr_row, column=kpi_col, value=str(kpi_label).upper())
            lbl_cell.font = Font(name=font_family, size=8, bold=True, color="64748B")
            lbl_cell.fill = zebra_fill
            lbl_cell.border = cell_border
            lbl_cell.alignment = Alignment(horizontal="center", vertical="center")

            val_cell = ws.cell(row=curr_row + 1, column=kpi_col, value=kpi_val)
            val_cell.font = Font(name=font_family, size=12, bold=True, color="0F172A")
            val_cell.fill = zebra_fill
            val_cell.border = cell_border
            val_cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(kpi_val, (int, float)):
                val_cell.number_format = "#,##0.00" if isinstance(kpi_val, float) else "#,##0"

            kpi_col += 1
        curr_row += 3

    # 3. Table Headers
    hdr_row_idx = curr_row
    for c_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=hdr_row_idx, column=c_idx, value=h)
        c.font = tbl_hdr_font
        c.fill = tbl_hdr_fill
        c.border = cell_border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hdr_row_idx].height = 24
    curr_row += 1

    # 4. Data Rows
    column_totals = {idx: 0.0 for idx in tot_indices}

    for r_idx, r in enumerate(rows):
        is_even = (r_idx % 2 == 1)
        row_fill = zebra_fill if is_even else PatternFill(fill_type=None)
        row_num = curr_row

        for c_idx, val in enumerate(r, start=1):
            cell = ws.cell(row=row_num, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = cell_border
            if row_fill.fill_type:
                cell.fill = row_fill

            # Check for status column highlighting
            str_val = str(val).lower().strip() if val is not None else ""
            if str_val in ("overdue", "rejected", "critical", "danger"):
                cell.fill = red_status_fill
                cell.font = red_status_font
                cell.alignment = Alignment(horizontal="center")
            elif str_val in ("paid", "active", "completed", "closed", "pass"):
                cell.fill = green_status_fill
                cell.font = green_status_font
                cell.alignment = Alignment(horizontal="center")
            elif str_val in ("pending", "partial", "in_progress", "awaiting_grn"):
                cell.fill = amber_status_fill
                cell.font = amber_status_font
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in curr_indices:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
                    if c_idx in tot_indices:
                        column_totals[c_idx] += float(val)
            elif c_idx in tot_indices:
                cell.alignment = Alignment(horizontal="right")
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0"
                    column_totals[c_idx] += float(val)
            elif isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, str) and (val.startswith("20") and len(val) >= 10 and "-" in val):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

        curr_row += 1

    # 5. Summary / Total Row
    if tot_indices and rows:
        tot_row_num = curr_row
        ws.cell(row=tot_row_num, column=1, value="TOTAL").font = Font(name=font_family, size=10, bold=True, color="0F172A")
        ws.cell(row=tot_row_num, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=tot_row_num, column=1).fill = total_fill
        ws.cell(row=tot_row_num, column=1).border = total_border

        for c_idx in range(2, len(headers) + 1):
            cell = ws.cell(row=tot_row_num, column=c_idx)
            cell.fill = total_fill
            cell.border = total_border
            if c_idx in tot_indices:
                cell.value = column_totals[c_idx]
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if c_idx in curr_indices:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"
            else:
                cell.value = ""

        ws.row_dimensions[tot_row_num].height = 20
        curr_row += 1

    # 6. Auto-fit Column Widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip title row and KPI cards when computing column width
            if cell.row < hdr_row_idx:
                continue
            val_str = str(cell.value or "")
            if cell.number_format and "#,##0" in cell.number_format and isinstance(cell.value, (int, float)):
                val_str = f"{cell.value:,.2f}"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_export_response(
    filename_base: str,
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    format: str = "xlsx",
    summary_kpis: Optional[Dict[str, Any]] = None,
    currency_col_indices: Optional[List[int]] = None,
    total_col_indices: Optional[List[int]] = None,
    sheet_name: str = "Data",
) -> StreamingResponse:
    """Builds a standardized StreamingResponse for either XLSX or CSV format."""
    fmt = (format or "xlsx").lower().strip()
    date_tag = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    if fmt == "csv":
        data_bytes = generate_csv_bytes(headers, rows)
        filename = f"{filename_base}_{date_tag}.csv"
        media_type = "text/csv; charset=utf-8"
    else:
        data_bytes = generate_excel_workbook(
            title=title,
            headers=headers,
            rows=rows,
            summary_kpis=summary_kpis,
            sheet_name=sheet_name,
            currency_col_indices=currency_col_indices,
            total_col_indices=total_col_indices,
        )
        filename = f"{filename_base}_{date_tag}.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    return StreamingResponse(
        io.BytesIO(data_bytes),
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )
