"""Daily Payment Upload Cadence & Reconciliation Verification Tests.

Covers:
1. Natural-key duplicate protection in daily payment import (re-upload skips duplicates).
2. Additive upload support (multiple daily files accumulate without wiping).
3. Progress tracking endpoint (Mon-Fri business days, MTD expected vs uploaded).
4. Configurable settlement lag setting (default 4 days).
5. 4-day settlement lag check (recent shipments marked pending_lag, not overdue).
6. UTR extraction and order_release_id fallback matching for unsettled rows.
7. Real data verification / math spot-check against downloaded Myntra fixtures.
"""

import os
import sys
import io
import pytest
import openpyxl
from datetime import datetime, timezone, timedelta

# Ensure backend root is in sys.path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "admin@sskfootcare.com")
ADMIN_PASS = os.environ.get("ADMIN_PASSWORD", "Admin@123")


@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient
    from server import app
    with TestClient(app, base_url="http://testserver/api") as tc:
        r = tc.post("/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASS})
        assert r.status_code == 200, f"Login failed: {r.text}"
        yield tc


class TestDailyPaymentUploadCadence:

    def test_settings_get_and_post(self, client):
        # 1. Get default settings
        res = client.get("/online-reconciliation/settings")
        assert res.status_code == 200
        data = res.json()
        assert "daily_payment_lag_days" in data
        assert "aged_pending_days" in data
        assert data["daily_payment_lag_days"] == 4
        assert data["aged_pending_days"] == 30

        # 2. Update settings
        res_post = client.post(
            "/online-reconciliation/settings",
            json={"daily_payment_lag_days": 5, "aged_pending_days": 45}
        )
        assert res_post.status_code == 200
        assert res_post.json()["daily_payment_lag_days"] == 5

        # Verify updated
        res_check = client.get("/online-reconciliation/settings")
        assert res_check.json()["daily_payment_lag_days"] == 5
        assert res_check.json()["aged_pending_days"] == 45

        # Reset back to default
        client.post(
            "/online-reconciliation/settings",
            json={"daily_payment_lag_days": 4, "aged_pending_days": 30}
        )

    def test_duplicate_protection_and_additive_upload(self, client):
        # Clear test collections
        r_clear = client.post("/online-reconciliation/clear-test-data")
        assert r_clear.status_code == 200

        # Day 1 CSV with 2 rows (Forward and Reverse on same order line to test natural key)
        csv_day1 = (
            "NEFT_Ref,Settled_Amount,Commission,Shipping_Fee,TDS,Payment_Type,Order_Type,order_release_id,seller_order_id,order_line_id,return_id,Payment_Date\n"
            "NEFT_D1,450.00,45.00,30.00,4.50,prepaid,Forward,REL_D1,SO_D1,LINE_D1,,2026-09-01\n"
            "NEFT_D1,-50.00,-5.00,0.00,0.00,prepaid,Reverse,REL_D1,SO_D1,LINE_D1,RET_D1,2026-09-01\n"
        )
        files1 = {"file": ("daily_2026-09-01.csv", csv_day1.encode("utf-8"), "text/csv")}
        res1 = client.post("/online-reconciliation/import-daily-payments", files=files1)
        assert res1.status_code == 200
        j1 = res1.json()
        assert j1["count"] == 2
        assert j1["inserted"] == 2
        assert j1["skipped_duplicates"] == 0

        # Re-upload identical CSV: duplicate protection must catch both rows!
        files1_dup = {"file": ("daily_2026-09-01_reupload.csv", csv_day1.encode("utf-8"), "text/csv")}
        res1_dup = client.post("/online-reconciliation/import-daily-payments", files=files1_dup)
        assert res1_dup.status_code == 200
        j1_dup = res1_dup.json()
        assert j1_dup["count"] == 0
        assert j1_dup["inserted"] == 0
        assert j1_dup["skipped_duplicates"] == 2
        assert j1_dup["total_in_file"] == 2

        # Day 2 CSV: 1 new row + 1 duplicate row from day 1
        csv_day2 = (
            "NEFT_Ref,Settled_Amount,Commission,Shipping_Fee,TDS,Payment_Type,Order_Type,order_release_id,seller_order_id,order_line_id,return_id,Payment_Date\n"
            "NEFT_D1,450.00,45.00,30.00,4.50,prepaid,Forward,REL_D1,SO_D1,LINE_D1,,2026-09-01\n"  # duplicate
            "NEFT_D2,600.00,60.00,35.00,6.00,postpaid,Forward,REL_D2,SO_D2,LINE_D2,,2026-09-02\n"  # new
        )
        files2 = {"file": ("daily_2026-09-02.csv", csv_day2.encode("utf-8"), "text/csv")}
        res2 = client.post("/online-reconciliation/import-daily-payments", files=files2)
        assert res2.status_code == 200
        j2 = res2.json()
        assert j2["inserted"] == 1
        assert j2["skipped_duplicates"] == 1
        assert j2["total_in_file"] == 2

    def test_progress_tracking_endpoint(self, client):
        # We uploaded dates 2026-09-01 (Tuesday) and 2026-09-02 (Wednesday)
        res = client.get("/online-reconciliation/daily-payments/progress?month=2026-09")
        assert res.status_code == 200
        data = res.json()
        assert data["month"] == "2026-09"
        assert data["total_business_days_in_month"] >= 20
        assert data["uploaded_business_days_count"] >= 2
        assert "2026-09-01" in data["uploaded_business_days"]
        assert "2026-09-02" in data["uploaded_business_days"]
        assert data["total_rows_this_month"] == 3  # 2 rows from Day 1 + 1 new from Day 2

    def test_settlement_lag_and_unsettled_fallback(self, client):
        # Clear test collections
        client.post("/online-reconciliation/clear-test-data")

        # 1. Monthly orders:
        # Order A: delivered 2 days ago (within 4-day lag), absent from settlements -> should be pending_lag
        # Order B: delivered 20 days ago (overdue), absent from settlements -> should be unmatched
        # Order C: delivered 10 days ago, in unsettled excel without UTR -> matches via order_release_id fallback!
        # Order D: delivered 10 days ago, in settled excel with UTR -> matches settled!
        today = datetime.now(timezone.utc).date()
        recent_date = (today - timedelta(days=2)).isoformat()
        overdue_date = (today - timedelta(days=20)).isoformat()
        settled_date = (today - timedelta(days=10)).isoformat()

        monthly_csv = (
            "seller order id,order release id,sku id,style id,seller sku code,size,order status,packed on,shipped on,delivered on,cancelled on,rto/return creation date,final amount,seller price\n"
            f"SO_RECENT,REL_RECENT,SKU_REC,STYLE-X,SSK-X-8,8,Delivered,{recent_date},{recent_date},{recent_date},,,500,500\n"
            f"SO_OVERDUE,REL_OVERDUE,SKU_OVD,STYLE-X,SSK-X-9,9,Delivered,{overdue_date},{overdue_date},{overdue_date},,,600,600\n"
            f"SO_UNSETTLED,REL_UNSETTLED,SKU_UNS,STYLE-X,SSK-X-7,7,Shipped,{settled_date},{settled_date},,,,400,400\n"
            f"SO_SETTLED,REL_SETTLED,SKU_SET,STYLE-X,SSK-X-8,8,Delivered,{settled_date},{settled_date},{settled_date},,,700,700\n"
        )
        r_m = client.post(
            "/online-reconciliation/import-monthly-report",
            files={"file": ("monthly.csv", monthly_csv.encode("utf-8"), "text/csv")}
        )
        assert r_m.status_code == 200

        # 2. Settled Excel with UTR_Number_Prepaid column
        wb_s = openpyxl.Workbook()
        ws_s = wb_s.active
        ws_s.title = "forward_settled"
        ws_s.append(["Row 1"])
        ws_s.append(["Row 2"])
        ws_s.append(["order_release_id", "seller_order_id", "sku_id", "style_id", "Settled_Amount_Postpaid", "Settled_Amount_Prepaid", "Commission_Amount_incl_GST", "Logistics_Cost_Forward_incl_Tax", "Fixed_Fee", "Pick_and_Pack_Fees", "Tech_Enablement_Charges", "UTR_Number_Postpaid", "UTR_Number_Prepaid"])
        # Empty string for postpaid UTR, valid value in prepaid UTR:
        ws_s.append(["REL_SETTLED", "SO_SETTLED", "SKU_SET", "STYLE-X", 0, 700.0, 70.0, 40.0, 10.0, 5.0, 2.0, "", "AXIS_PREPAID_12345"])
        buf_s = io.BytesIO()
        wb_s.save(buf_s)
        r_s = client.post(
            "/online-reconciliation/import-settlements",
            files={"file": ("settled.xlsx", buf_s.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        assert r_s.status_code == 200

        # 3. Unsettled Excel: lacks seller_order_id, has order_release_id and NO UTR!
        wb_u = openpyxl.Workbook()
        ws_u = wb_u.active
        ws_u.title = "forward_unsettled"
        ws_u.append(["Row 1"])
        ws_u.append(["Row 2"])
        ws_u.append(["order_release_id", "seller_order_id", "sku_id", "style_id", "Amount_pending_settlement_Postpaid", "Amount_pending_settlement_Prepaid"])
        ws_u.append(["REL_UNSETTLED", "", "", "STYLE-X", 400.0, 0])
        buf_u = io.BytesIO()
        wb_u.save(buf_u)
        r_u = client.post(
            "/online-reconciliation/import-settlements",
            files={"file": ("unsettled.xlsx", buf_u.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        )
        assert r_u.status_code == 200

        # 4. Run reconciliation with daily_payment_lag_days=4
        r_run = client.post("/online-reconciliation/run?daily_payment_lag_days=4")
        assert r_run.status_code == 200
        summary = r_run.json()

        # SO_SETTLED should be settled
        assert summary["settled_count"] >= 1

        # SO_UNSETTLED should match via order_release_id fallback!
        assert summary["pending_count"] >= 1

        # SO_RECENT (delivered 2 days ago <= 4 days lag) should be pending_lag
        assert summary["pending_lag_count"] >= 1

        # SO_OVERDUE (delivered 20 days ago > 4 days lag) should be unmatched
        assert summary["unmatched_count"] >= 1

        # Check unreconciled list contains explicit reason with lag note
        unrec = summary["unreconciled_orders"]
        recent_entry = next((u for u in unrec if u["seller_order_id"] == "SO_RECENT"), None)
        assert recent_entry is not None
        assert recent_entry["status"] == "pending_lag"

        overdue_entry = next((u for u in unrec if u["seller_order_id"] == "SO_OVERDUE"), None)
        assert overdue_entry is not None
        assert "Absent from settlement files" in overdue_entry["reasons"]

    def test_spot_check_real_fixtures(self, client):
        import glob
        download_dir = "C:/Users/Dell/Downloads"
        prepaid_files = glob.glob(f"{download_dir}/part-*11465*.csv")
        postpaid_files = glob.glob(f"{download_dir}/part-*11501*.csv")

        if not prepaid_files or not postpaid_files:
            pytest.skip("Downloaded Myntra fixture files not available in C:/Users/Dell/Downloads")

        client.post("/online-reconciliation/clear-test-data")

        # Import prepaid daily file
        with open(prepaid_files[0], "rb") as f:
            r1 = client.post("/online-reconciliation/import-daily-payments", files={"file": ("prepaid.csv", f.read(), "text/csv")})
            assert r1.status_code == 200
            j1 = r1.json()
            assert j1["inserted"] == 39
            assert j1["skipped_duplicates"] == 0

        # Re-upload prepaid daily file: must skip all 39 duplicates!
        with open(prepaid_files[0], "rb") as f:
            r1_re = client.post("/online-reconciliation/import-daily-payments", files={"file": ("prepaid_re.csv", f.read(), "text/csv")})
            assert r1_re.status_code == 200
            j1_re = r1_re.json()
            assert j1_re["inserted"] == 0
            assert j1_re["skipped_duplicates"] == 39

        # Import postpaid daily file: 25 rows inserted, additive (total 64)
        with open(postpaid_files[0], "rb") as f:
            r2 = client.post("/online-reconciliation/import-daily-payments", files={"file": ("postpaid.csv", f.read(), "text/csv")})
            assert r2.status_code == 200
            j2 = r2.json()
            assert j2["inserted"] == 25
            assert j2["skipped_duplicates"] == 0

        # Check progress endpoint for Sep 2026 (dates in these fixtures are 2026-09-02)
        r_prog = client.get("/online-reconciliation/daily-payments/progress?month=2026-09")
        assert r_prog.status_code == 200
        prog = r_prog.json()
        assert "2026-09-02" in prog["uploaded_business_days"]
        assert prog["total_rows_this_month"] == 64

    def test_unsettled_fallback_amount_and_date_and_no_utr(self, client):
        """Verifies:
        1. Unsettled rows never have UTRs, and matching succeeds without UTR errors.
        2. Unsettled matching uses order_release_id + amount + date fallback.
        3. Settled rows match UTR against daily payments.
        4. Cadence distinction: monthly settlement import does not wipe continuous daily payments.
        """
        client.post("/online-reconciliation/clear-test-data")

        # 1. Pre-load continuous daily payment
        daily_csv = (
            "NEFT_Ref,Settled_Amount,Commission,Shipping_Fee,TDS,Payment_Type,Order_Type,order_release_id,seller_order_id,order_line_id,return_id,Payment_Date\n"
            "NEFT_SETTLED_999,650.00,65.00,30.00,6.50,prepaid,Forward,REL_SETTLED_1,SO_SETTLED_1,LINE_SET_1,,2026-08-15\n"
        )
        r_dp = client.post("/online-reconciliation/import-daily-payments", files={"file": ("daily.csv", daily_csv.encode("utf-8"), "text/csv")})
        assert r_dp.status_code == 200
        assert r_dp.json()["inserted"] == 1

        # 2. Monthly Order Report:
        # Order 1: Delivered, will match settled file with UTR
        # Order 2: Shipped, will match unsettled file without UTR via fallback (rel_id + amount + date)
        today = datetime.now(timezone.utc).date()
        date_15d_ago = (today - timedelta(days=15)).isoformat()
        monthly_csv = (
            "seller order id,order release id,sku id,style id,seller sku code,size,order status,packed on,shipped on,delivered on,cancelled on,rto/return creation date,final amount,seller price\n"
            f"SO_SETTLED_1,REL_SETTLED_1,SKU_SET,STYLE-A,SSK-A-8,8,Delivered,{date_15d_ago},{date_15d_ago},{date_15d_ago},,,750,750\n"
            f"SO_UNKNOWN_FB,REL_UNSETTLED_FB,SKU_MISMATCH,STYLE-A,SSK-A-8,8,Shipped,{date_15d_ago},{date_15d_ago},,,,480,480\n"
        )
        r_m = client.post("/online-reconciliation/import-monthly-report", files={"file": ("monthly.csv", monthly_csv.encode("utf-8"), "text/csv")})
        assert r_m.status_code == 200

        # 3. Settled file with UTR
        wb_s = openpyxl.Workbook()
        ws_s = wb_s.active
        ws_s.title = "forward_settled"
        ws_s.append(["Row 1"])
        ws_s.append(["Row 2"])
        ws_s.append(["order_release_id", "seller_order_id", "sku_id", "style_id", "Settled_Amount_Postpaid", "Settled_Amount_Prepaid", "Commission_Amount_incl_GST", "Logistics_Cost_Forward_incl_Tax", "Fixed_Fee", "Pick_and_Pack_Fees", "Tech_Enablement_Charges", "UTR_Number_Postpaid", "UTR_Number_Prepaid"])
        ws_s.append(["REL_SETTLED_1", "SO_SETTLED_1", "SKU_SET", "STYLE-A", 0, 650.0, 65.0, 30.0, 5.0, 0, 0, "", "NEFT_SETTLED_999"])
        buf_s = io.BytesIO()
        wb_s.save(buf_s)
        r_s = client.post("/online-reconciliation/import-settlements", files={"file": ("settled.xlsx", buf_s.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r_s.status_code == 200

        # 4. Unsettled file: 2 candidates with same release ID (REL_UNSETTLED_FB) but NO UTR!
        # Candidate 1: amount=480.0, date matches date_15d_ago -> Should match Order 2 via amount+date fallback!
        # Candidate 2: amount=120.0, date=2025-01-01 -> Non-matching amount/date
        wb_u = openpyxl.Workbook()
        ws_u = wb_u.active
        ws_u.title = "forward_unsettled"
        ws_u.append(["Row 1"])
        ws_u.append(["Row 2"])
        ws_u.append(["order_release_id", "seller_order_id", "sku_id", "style_id", "Amount_pending_settlement_Postpaid", "Amount_pending_settlement_Prepaid", "order_date", "UTR_Number_Postpaid", "UTR_Number_Prepaid"])
        ws_u.append(["REL_UNSETTLED_FB", "", "", "STYLE-A", 120.0, 0, "2025-01-01", "", ""])
        ws_u.append(["REL_UNSETTLED_FB", "", "", "STYLE-A", 480.0, 0, date_15d_ago, "", ""])
        buf_u = io.BytesIO()
        wb_u.save(buf_u)
        r_u = client.post("/online-reconciliation/import-settlements", files={"file": ("unsettled.xlsx", buf_u.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r_u.status_code == 200

        # 5. Run reconciliation
        r_run = client.post("/online-reconciliation/run?daily_payment_lag_days=4&aged_pending_days=30")
        assert r_run.status_code == 200
        summary = r_run.json()

        # Both orders must be reconciled (100% join rate)
        assert summary["join_rate_pct"] == 100.0
        assert summary["settled_count"] == 1
        # Unsettled order matched via fallback without UTR:
        assert summary["pending_count"] == 1
        assert summary["unmatched_count"] == 0

        # Settled UTR matched daily payments with 0 mismatch diff:
        assert len(summary["neft_mismatches"]) == 0

