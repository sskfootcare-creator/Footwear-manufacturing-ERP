"""Tests for Dispatch and Merged Packing List generation:
1. PO style code (external_sku / mapped_from_sku / po_style_code) used in Style column.
2. Single consolidated row per style & color combination (carton ranges e.g. 1-10).
3. Populated size matrix across size columns.
4. Elimination of unnecessary padded empty rows (Grand Total follows data immediately).
5. Dynamic size detection (e.g. UK sizes 4-8 vs default 36-42).
"""

import io
import openpyxl
import pytest
from unittest.mock import MagicMock, AsyncMock
from bson import ObjectId

from packing_list import build_dispatch_packing_list, _carton_po_style, _resolve_sizes
from routes.invoice_packing import _enrich_cartons_with_mapped_sku


def test_carton_po_style_priority():
    """Verify priority order: po_style_code > mapped_from_sku > external_sku > customer_style_code > po line items > style_code."""
    # 1. Direct po_style_code
    c1 = {"po_style_code": "PO-100", "mapped_from_sku": "MAP-100", "style_code": "SSK-001"}
    assert _carton_po_style(c1) == "PO-100"

    # 2. mapped_from_sku
    c2 = {"mapped_from_sku": "MAP-200", "style_code": "SSK-002"}
    assert _carton_po_style(c2) == "MAP-200"

    # 3. external_sku
    c3 = {"external_sku": "EXT-300", "style_code": "SSK-003"}
    assert _carton_po_style(c3) == "EXT-300"

    # 4. Fallback to PO line items
    po = {
        "line_items": [
            {"style_code": "SSK-004", "color": "Black", "size": "40", "external_sku": "PO-EXT-400"}
        ]
    }
    c4 = {"style_code": "SSK-004", "color": "Black", "size": "40"}
    assert _carton_po_style(c4, po) == "PO-EXT-400"

    # 5. Fallback to carton style_code
    c5 = {"style_code": "SSK-005", "color": "Brown"}
    assert _carton_po_style(c5) == "SSK-005"


def test_dispatch_packing_list_consolidated_and_no_padded_rows():
    """Verify build_dispatch_packing_list produces:
    - Exactly 1 row per (style, color)
    - PO style codes
    - Populated size matrix
    - Grand Total immediately following data rows (no 23 padded rows)
    """
    po = {
        "po_number": "PO-TEST-001",
        "po_date": "10/09/2026",
        "client_name": "TEST BUYER",
        "line_items": [
            {"style_code": "SSK_00004", "color": "Black", "size": "38", "quantity": 60, "external_sku": "BUYER-STYLE-101"},
            {"style_code": "SSK_00004", "color": "Black", "size": "39", "quantity": 80, "external_sku": "BUYER-STYLE-101"},
            {"style_code": "SSK_00004", "color": "Black", "size": "40", "quantity": 60, "external_sku": "BUYER-STYLE-101"},
            {"style_code": "SSK_00007", "color": "Tan", "size": "40", "quantity": 40, "mapped_from_sku": "BUYER-STYLE-202"},
        ]
    }

    # 10 cartons for Style 1 Black (3 of size 38, 4 of size 39, 3 of size 40)
    # 2 cartons for Style 2 Tan (2 of size 40)
    cartons = []
    box = 1
    for _ in range(3):
        cartons.append({"box_number": box, "style_code": "SSK_00004", "po_style_code": "BUYER-STYLE-101", "color": "Black", "size": "38", "qty": 20})
        box += 1
    for _ in range(4):
        cartons.append({"box_number": box, "style_code": "SSK_00004", "po_style_code": "BUYER-STYLE-101", "color": "Black", "size": "39", "qty": 20})
        box += 1
    for _ in range(3):
        cartons.append({"box_number": box, "style_code": "SSK_00004", "po_style_code": "BUYER-STYLE-101", "color": "Black", "size": "40", "qty": 20})
        box += 1
    for _ in range(2):
        cartons.append({"box_number": box, "style_code": "SSK_00007", "mapped_from_sku": "BUYER-STYLE-202", "color": "Tan", "size": "40", "qty": 20})
        box += 1

    xlsx_bytes = build_dispatch_packing_list(cartons, po, invoice_no="INV-2026-001")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    ws = wb.active

    # Data starts at row 11
    # Style 1 row should be row 11
    assert ws["B11"].value == "BUYER-STYLE-101"
    assert ws["C11"].value == "Black"
    assert ws["D11"].value == "1-10"
    # Size 38 is Col G (cols: A=SITE, B=Style, C=Color, D=CTN.NO, E=36, F=37, G=38, H=39, I=40, J=41, K=42)
    assert ws["G11"].value == 60   # Size 38 qty
    assert ws["H11"].value == 80   # Size 39 qty
    assert ws["I11"].value == 60   # Size 40 qty
    assert ws["L11"].value == 200  # PCS/CTN
    assert ws["M11"].value == 20   # Per Carton
    assert ws["N11"].value == 10   # TTL CTN
    assert ws["O11"].value == 200  # Total PCS

    # Style 2 row should be row 12
    assert ws["B12"].value == "BUYER-STYLE-202"
    assert ws["C12"].value == "Tan"
    assert ws["D12"].value == "11-12"
    assert ws["I12"].value == 40   # Size 40 qty
    assert ws["N12"].value == 2    # TTL CTN
    assert ws["O12"].value == 40   # Total PCS

    # Grand Total row MUST be row 13 immediately following row 12 (NO padded rows!)
    assert ws["A13"].value == "GRAND TOTAL"
    assert ws["G13"].value == "=SUM(G11:G12)"
    assert ws["H13"].value == "=SUM(H11:H12)"
    assert ws["I13"].value == "=SUM(I11:I12)"
    assert ws["N13"].value == "=SUM(N11:N12)"
    assert ws["O13"].value == "=SUM(O11:O12)"

    # Row 14, 15 are empty spacer rows
    # Row 16 is Order Summary
    assert ws["B16"].value == "ORDER SUMMERY"


def test_dispatch_packing_list_dynamic_sizes():
    """Verify build_dispatch_packing_list dynamically detects and formats custom size sets (e.g. UK sizes 4-8)."""
    po = {
        "po_number": "PO-SIYARAM-001",
        "po_date": "10/09/2026",
        "client_name": "SIYARAM",
        "line_items": [
            {"style_code": "SSK_UK_1", "color": "Brown", "size": "4", "quantity": 20, "external_sku": "SIY-UK-01"},
            {"style_code": "SSK_UK_1", "color": "Brown", "size": "5", "quantity": 40, "external_sku": "SIY-UK-01"},
            {"style_code": "SSK_UK_1", "color": "Brown", "size": "6", "quantity": 60, "external_sku": "SIY-UK-01"},
            {"style_code": "SSK_UK_1", "color": "Brown", "size": "7", "quantity": 40, "external_sku": "SIY-UK-01"},
            {"style_code": "SSK_UK_1", "color": "Brown", "size": "8", "quantity": 20, "external_sku": "SIY-UK-01"},
        ]
    }
    cartons = [
        {"box_number": 1, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "4", "qty": 20},
        {"box_number": 2, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "5", "qty": 20},
        {"box_number": 3, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "5", "qty": 20},
        {"box_number": 4, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "6", "qty": 20},
        {"box_number": 5, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "6", "qty": 20},
        {"box_number": 6, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "6", "qty": 20},
        {"box_number": 7, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "7", "qty": 20},
        {"box_number": 8, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "7", "qty": 20},
        {"box_number": 9, "style_code": "SSK_UK_1", "external_sku": "SIY-UK-01", "color": "Brown", "size": "8", "qty": 20},
    ]

    xlsx_bytes = build_dispatch_packing_list(cartons, po, invoice_no="INV-UK-001")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    ws = wb.active

    # Check header size columns: E10, F10, G10, H10, I10 should be 4, 5, 6, 7, 8
    assert ws["E10"].value == "4"
    assert ws["F10"].value == "5"
    assert ws["G10"].value == "6"
    assert ws["H10"].value == "7"
    assert ws["I10"].value == "8"

    # Only 1 data row at row 11
    assert ws["B11"].value == "SIY-UK-01"
    assert ws["C11"].value == "Brown"
    assert ws["D11"].value == "1-9"
    assert ws["E11"].value == 20  # Size 4
    assert ws["F11"].value == 40  # Size 5
    assert ws["G11"].value == 60  # Size 6
    assert ws["H11"].value == 40  # Size 7
    assert ws["I11"].value == 20  # Size 8

    # Grand Total at row 12
    assert ws["A12"].value == "GRAND TOTAL"


@pytest.mark.anyio
async def test_enrich_cartons_with_mapped_sku_fallback():
    """Verify _enrich_cartons_with_mapped_sku correctly decorates cartons using PO line items."""
    mock_db = MagicMock()
    mock_db.production_jobs.find.return_value.to_list = AsyncMock(return_value=[])
    mock_db.sku_map.find.return_value.to_list = AsyncMock(return_value=[])

    po_id = ObjectId()
    mock_po = {
        "_id": po_id,
        "line_items": [
            {"style_code": "SSK_00004", "color": "Black", "size": "38", "mapped_from_sku": "BUYER-SKU-99"},
        ]
    }
    mock_db.pos.find.return_value.to_list = AsyncMock(return_value=[mock_po])

    cartons = [
        {"po_id": str(po_id), "style_code": "SSK_00004", "color": "Black", "size": "38", "qty": 20},
        # Color match without size match should also resolve style code
        {"po_id": str(po_id), "style_code": "SSK_00004", "color": "Black", "size": "39", "qty": 20},
    ]

    enriched = await _enrich_cartons_with_mapped_sku(cartons, db=mock_db)
    assert enriched[0].get("po_style_code") == "BUYER-SKU-99"
    assert enriched[1].get("po_style_code") == "BUYER-SKU-99"
