import pytest
import io
import os
import openpyxl
from httpx import AsyncClient, ASGITransport
from motor.motor_asyncio import AsyncIOMotorClient
import server
from server import app
from auth import create_access_token

TEST_STYLES = [
    {
        "code": "S-ZEC-001",
        "name": "ZECODE T-Strap Flat Sandal 01",
        "brand": "ZECODE",
        "category": "Sandals",
        "color": "Black",
        "upper_material": "Synthetic Leather",
        "insole_lining": "Comfort Cushion / Synthetic",
        "outsole": "TPR / Sheet Sole",
        "heel_type": "Flat (0.5 inch)",
        "transfer_price": 550.0,
        "target_mrp": 1299.0,
        "cost": 550.0,
        "image_url": "/api/uploads/images/catalogue/S-ZEC-001.png",
        "image_thumbnail_url": "/api/uploads/images/catalogue/S-ZEC-001.png",
        "bom": [{"item": "Synthetic", "rate": 200.0, "quantity": 1.0}],
        "labor": [{"role": "stitching", "rate": 50.0, "rate_per_pair": 50.0}],
        "packing_cost": 20.0,
    },
    {
        "code": "S-ZEC-002",
        "name": "ZECODE T-Strap Flat Sandal 02",
        "brand": "ZECODE",
        "category": "Sandals",
        "color": "Tan",
        "upper_material": "Synthetic Suede",
        "insole_lining": "Padded Insole / Rexine",
        "outsole": "TPR / Sheet Sole",
        "heel_type": "Flat (0.5 inch)",
        "transfer_price": 550.0,
        "target_mrp": 1299.0,
        "cost": 550.0,
        "image_url": "/api/uploads/images/catalogue/S-ZEC-002.png",
        "image_thumbnail_url": "/api/uploads/images/catalogue/S-ZEC-002.png",
        "bom": [{"item": "Synthetic", "rate": 220.0, "quantity": 1.0}],
        "labor": [{"role": "stitching", "rate": 50.0, "rate_per_pair": 50.0}],
        "packing_cost": 20.0,
    },
    {
        "code": "S-ZEC-003",
        "name": "ZECODE T-Strap Flat Sandal 03",
        "brand": "ZECODE",
        "category": "Sandals",
        "color": "Gold",
        "upper_material": "Synthetic Metallic",
        "insole_lining": "Cushioned PU",
        "outsole": "TPR / Sheet Sole",
        "heel_type": "Flat (0.5 inch)",
        "transfer_price": 550.0,
        "target_mrp": 1299.0,
        "cost": 550.0,
        "image_url": "/api/uploads/images/catalogue/S-ZEC-003.png",
        "image_thumbnail_url": "/api/uploads/images/catalogue/S-ZEC-003.png",
        "bom": [{"item": "Synthetic", "rate": 230.0, "quantity": 1.0}],
        "labor": [{"role": "stitching", "rate": 50.0, "rate_per_pair": 50.0}],
        "packing_cost": 20.0,
    },
]

@pytest.fixture(autouse=True)
def setup_db_loop():
    mongo_url = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
    db_name = os.environ.get("DB_NAME", "ssk_footwear_erp")
    server.client = AsyncIOMotorClient(mongo_url)
    server.db = server.client[db_name]
    yield

async def get_admin_headers():
    # 1. Ensure test style images exist on disk for CI / clean checkout environments
    curr_dir = os.path.dirname(os.path.abspath(__file__))  # backend/tests
    backend_dir = os.path.dirname(curr_dir)                # backend
    repo_root = os.path.dirname(backend_dir)               # root
    for base in [backend_dir, repo_root]:
        cat_dir = os.path.join(base, "uploads", "images", "catalogue")
        os.makedirs(cat_dir, exist_ok=True)
        for code in ["S-ZEC-001", "S-ZEC-002", "S-ZEC-003"]:
            target_file = os.path.join(cat_dir, f"{code}.png")
            if not os.path.isfile(target_file):
                from PIL import Image as PILImage
                img = PILImage.new("RGB", (100, 100), color=(73, 109, 137))
                img.save(target_file)

    # 2. Ensure test styles exist in db.styles for any test runner environment
    for s in TEST_STYLES:
        await server.db.styles.update_one(
            {"code": s["code"]},
            {"$set": s},
            upsert=True,
        )

    admin_user = await server.db.users.find_one({"role": "admin"})
    if not admin_user:
        admin_user = await server.db.users.find_one()
    if not admin_user:
        from auth import hash_password
        admin_doc = {
            "email": "admin@sskfootcare.com",
            "password_hash": hash_password("Admin@123"),
            "role": "admin",
            "name": "Admin User",
            "active": True,
        }
        res = await server.db.users.insert_one(admin_doc)
        admin_user = admin_doc
        admin_user["_id"] = res.inserted_id
    admin_id = str(admin_user["_id"])
    admin_email = admin_user.get("email", "admin@sskfootwear.com")
    token = create_access_token(admin_id, admin_email, "admin")
    return {"Authorization": f"Bearer {token}"}

@pytest.mark.anyio
async def test_catalogue_download_with_pictures():
    headers = await get_admin_headers()
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Test downloading catalogue with style codes
        res = await client.post(
            "/api/styles/catalogue/download",
            json={
                "style_codes": ["S-ZEC-001", "S-ZEC-002"],
                "brand": "ZECODE",
                "include_images": True,
            },
            headers=headers
        )
        assert res.status_code == 200, res.text
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
        
        # Verify valid excel file content
        content = res.content
        wb = openpyxl.load_workbook(io.BytesIO(content))
        assert "Footwear Costing & Specs" in wb.sheetnames
        ws = wb["Footwear Costing & Specs"]
        assert "SSK FOOTCARE - MASTER PRODUCT SPECIFICATION & COSTING SHEET" in ws["A1"].value
        assert "ZECODE" in ws["A2"].value
        assert ws["A4"].value == "Sr. No."
        assert ws["B4"].value == "Photo"
        assert ws["C4"].value == "Vendor SKU"
        assert ws["C5"].value == "S-ZEC-001"
        assert ws["C6"].value == "S-ZEC-002"
        # Check that images are present in sheet
        assert len(ws._images) >= 1

        # Test GET streaming endpoint as well
        get_res = await client.get(
            "/api/styles/catalogue/download?brand=ZECODE&include_images=true",
            headers=headers
        )
        assert get_res.status_code == 200


@pytest.mark.anyio
async def test_po_line_items_have_style_photo():
    headers = await get_admin_headers()
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        po_num = "PO-TEST-PHOTO-001"
        await server.db.pos.delete_many({"po_number": po_num})

        # Create a PO with style S-ZEC-001
        po_payload = {
            "po_number": po_num,
            "po_date": "2026-10-29",
            "client_name": "Test Client Fashion",
            "client_address": "123 Fashion Blvd",
            "delivery_date": "2026-11-20",
            "line_items": [
                {
                    "style_code": "S-ZEC-001",
                    "description": "T-Strap Flat",
                    "color": "Black",
                    "size": "7",
                    "quantity": 100,
                    "unit_price": 750.0,
                    "amount": 75000.0,
                }
            ],
            "subtotal": 75000.0,
            "cgst_rate": 2.5,
            "sgst_rate": 2.5,
            "igst_rate": 0,
        }
        res = await client.post(
            "/api/pos",
            json=po_payload,
            headers=headers
        )
        assert res.status_code in (200, 201), res.text
        created_po = res.json()
        po_id = created_po["id"]

        # Fetch PO detail
        get_res = await client.get(
            f"/api/pos/{po_id}",
            headers=headers
        )
        assert get_res.status_code == 200
        po_detail = get_res.json()
        line = po_detail["line_items"][0]
        # Check style photo attached to line item
        assert "image_url" in line or "image_thumbnail_url" in line
        img_url = line.get("image_url") or line.get("image_thumbnail_url")
        assert "S-ZEC-001" in img_url or "catalogue" in img_url or bool(img_url)

        # Cleanup
        await client.delete(f"/api/pos/{po_id}", headers=headers)


@pytest.mark.anyio
async def test_multi_po_budget_calculation_and_alignment():
    headers = await get_admin_headers()
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        po1_num = "PO-BUDGET-TEST-1"
        po2_num = "PO-BUDGET-TEST-2"
        await server.db.pos.delete_many({"po_number": {"$in": [po1_num, po2_num]}})

        # Create two test POs
        po1_res = await client.post(
            "/api/pos",
            json={
                "po_number": po1_num,
                "po_date": "2026-10-29",
                "client_name": "Client A",
                "client_address": "Street 1",
                "delivery_date": "2026-11-20",
                "line_items": [
                    {
                        "style_code": "S-ZEC-001",
                        "description": "Style 1",
                        "color": "Black",
                        "size": "7",
                        "quantity": 50,
                        "unit_price": 800.0,
                        "amount": 40000.0,
                    },
                    {
                        "style_code": "S-ZEC-002",
                        "description": "Style 2",
                        "color": "Tan",
                        "size": "8",
                        "quantity": 30,
                        "unit_price": 850.0,
                        "amount": 25500.0,
                    }
                ],
                "subtotal": 65500.0,
            },
            headers=headers
        )
        assert po1_res.status_code in (200, 201), po1_res.text
        po1 = po1_res.json()

        po2_res = await client.post(
            "/api/pos",
            json={
                "po_number": po2_num,
                "po_date": "2026-10-29",
                "client_name": "Client B",
                "client_address": "Street 2",
                "delivery_date": "2026-11-20",
                "line_items": [
                    {
                        "style_code": "S-ZEC-003",
                        "description": "Style 3",
                        "color": "Gold",
                        "size": "6",
                        "quantity": 40,
                        "unit_price": 900.0,
                        "amount": 36000.0,
                    }
                ],
                "subtotal": 36000.0,
            },
            headers=headers
        )
        assert po2_res.status_code in (200, 201), po2_res.text
        po2 = po2_res.json()

        # 1. Test calculation for subset of styles (only S-ZEC-001 and S-ZEC-003, leaving out S-ZEC-002)
        calc_payload = {
            "items": [
                {
                    "po_id": po1["id"],
                    "po_number": po1["po_number"],
                    "client_name": po1["client_name"],
                    "style_code": "S-ZEC-001",
                    "quantity": 50,
                    "unit_price": 800.0,
                },
                {
                    "po_id": po2["id"],
                    "po_number": po2["po_number"],
                    "client_name": po2["client_name"],
                    "style_code": "S-ZEC-003",
                    "quantity": 40,
                    "unit_price": 900.0,
                }
            ],
            "po_ids": [po1["id"], po2["id"]],
        }

        calc_res = await client.post(
            "/api/pos/multi-budget/calculate",
            json=calc_payload,
            headers=headers
        )
        assert calc_res.status_code == 200, calc_res.text
        calc_data = calc_res.json()

        # Expected total quantity: 50 + 40 = 90
        assert calc_data["total_quantity"] == 90
        # Expected revenue: 50*800 + 40*900 = 40000 + 36000 = 76000
        assert calc_data["total_revenue"] == 76000.0
        assert calc_data["total_production_budget"] > 0
        assert "total_bom_cost" in calc_data
        assert "total_labor_cost" in calc_data
        assert "total_packing_cost" in calc_data
        assert len(calc_data["line_items"]) == 2

        # 2. Test saving official alignment snapshot
        align_res = await client.post(
            "/api/pos/multi-budget/align",
            json={
                "items": calc_payload["items"],
                "po_ids": [po1["id"], po2["id"]],
                "notes": "Planning alignment test",
            },
            headers=headers
        )
        assert align_res.status_code == 200, align_res.text
        align_data = align_res.json()
        assert "id" in align_data
        assert align_data["summary"]["total_quantity"] == 90
        assert align_data["summary"]["total_revenue"] == 76000.0

        # Cleanup
        await client.delete(f"/api/pos/{po1['id']}", headers=headers)
        await client.delete(f"/api/pos/{po2['id']}", headers=headers)
