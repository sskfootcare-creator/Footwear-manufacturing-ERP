import os
import sys
import openpyxl
import re
import asyncio
from motor.motor_asyncio import AsyncIOMotorClient

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from routes.online_returns_engine import get_footwear_placeholder_image

SAMPLE_PNL_PATH = r"C:\Users\Dell\Downloads\PnLReport_37353.xlsx"

async def run_backfill():
    style_to_myntra_id = {}
    if os.path.exists(SAMPLE_PNL_PATH):
        wb = openpyxl.load_workbook(SAMPLE_PNL_PATH, data_only=True)
        if "SKU_Detail" in wb.sheetnames:
            ws = wb["SKU_Detail"]
            header = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
            if "sku_code" in header and "style_id" in header:
                sku_idx = header.index("sku_code") + 1
                style_id_idx = header.index("style_id") + 1
                for r in range(2, ws.max_row + 1):
                    sku = str(ws.cell(r, sku_idx).value or "").strip()
                    sid = str(ws.cell(r, style_id_idx).value or "").strip()
                    m = re.match(r"^([A-Za-z0-9_\-]+?)[-_]([A-Za-z0-9]+)[-_]([0-9]+(?:\.[0-9]+)?)$", sku)
                    root = m.group(1) if m else sku
                    if sid:
                        style_to_myntra_id[root] = sid
                        style_to_myntra_id[root.replace("-", "_")] = sid
                        style_to_myntra_id[root.replace("_", "-")] = sid

    client = AsyncIOMotorClient("mongodb://localhost:27017")
    db = client.ssk_footwear_erp

    sku_maps = await db.sku_map.find({}).to_list(2000)
    map_by_ext = {}
    for sm in sku_maps:
        ext = str(sm.get("external_sku") or "").strip().lower()
        if ext:
            target_erp = sm.get("internal_style_code") or sm.get("style_code")
            map_by_ext[ext] = target_erp
            m = re.match(r"^([A-Za-z0-9_\-]+?)[-_]([A-Za-z0-9]+)[-_]([0-9]+(?:\.[0-9]+)?)$", ext)
            if m:
                map_by_ext[m.group(1).lower()] = target_erp

    erp_styles = await db.styles.find({}).to_list(2000)
    style_codes = {str(s.get("code") or "").strip().lower(): s for s in erp_styles if s.get("code")}

    custom_photos = {}
    async for cp in db.online_style_photos.find({}):
        custom_photos[cp["style_code"]] = cp["image_url"]

    cursor = db.online_monthly_reconciliation_overviews.find({})
    overviews = await cursor.to_list(100)
    for ov in overviews:
        updated_styles = []
        for s in ov.get("styles", []):
            st = s.get("style_code")
            myntra_id = (
                style_to_myntra_id.get(st)
                or style_to_myntra_id.get(st.replace("-", "_"))
                or style_to_myntra_id.get(st.replace("_", "-"))
                or s.get("myntra_style_id")
                or ""
            )

            st_low = st.lower()
            erp_code = (
                map_by_ext.get(st_low)
                or map_by_ext.get(st_low.replace("-", "_"))
                or map_by_ext.get(st_low.replace("_", "-"))
                or (style_codes[st_low].get("code") if st_low in style_codes else None)
                or s.get("erp_style_code")
                or ""
            )

            img = (
                custom_photos.get(st)
                or (style_codes[st_low].get("image_url") if st_low in style_codes else None)
                or s.get("image_url")
                or get_footwear_placeholder_image(st)
            )

            s["myntra_style_id"] = myntra_id
            s["erp_style_code"] = erp_code
            s["image_url"] = img
            updated_styles.append(s)

        await db.online_monthly_reconciliation_overviews.update_one(
            {"_id": ov["_id"]},
            {"$set": {"styles": updated_styles}}
        )
        print(f"Updated overview {ov.get('month')} ({ov.get('platform')}): {len(updated_styles)} styles.")

if __name__ == "__main__":
    asyncio.run(run_backfill())
